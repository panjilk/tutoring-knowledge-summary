"""将不同格式的本地文件解析为 LangChain 文档。"""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import fitz
from docx import Document as WordDocument
from langchain_core.documents import Document

from src.exceptions import DocumentParseError, UnsupportedFileError

# 上传层用此集合做白名单校验；新增格式时需要同时增加对应的解析分支。
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
# 匹配“第 N 章 + 标题”，用于把 PDF 页面归属到章节，支持常见中文数字。
CHAPTER_PATTERN = re.compile(
    r"第\s*([0-9一二三四五六七八九十百]+)\s*章\s*([^\n]{1,30})"
)
# 同一章页眉相隔不超过四页，视作同一段连续章节；用来排除目录中的偶然匹配。
MAX_CHAPTER_PAGE_GAP = 4


def load_document(
    file_path: Path,
    *,
    document_id: str,
    knowledge_base_id: str,
    original_name: str,
) -> list[Document]:
    """根据扩展名解析文件，同时保留引用所需的位置信息。"""

    # 先按扩展名路由到解析器。所有解析器都输出统一的 LangChain Document 类型。
    extension = file_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"暂不支持 {extension} 格式。")

    # 这些元数据会随每个文本单元、每个文档块进入向量库，用于来源引用和过滤检索。
    base_metadata = {
        "document_id": document_id,
        "knowledge_base_id": knowledge_base_id,
        "source_file": original_name,
        "file_type": extension.removeprefix("."),
    }

    try:
        if extension == ".pdf":
            documents = _load_pdf(file_path, base_metadata)
        elif extension == ".docx":
            documents = _load_docx(file_path, base_metadata)
        elif extension in {".md", ".markdown"}:
            documents = _load_markdown(file_path, base_metadata)
        else:
            documents = _load_text(file_path, base_metadata)
    except (UnsupportedFileError, DocumentParseError):
        raise
    except Exception as exc:
        raise DocumentParseError(
            f"解析文件 {original_name} 时发生错误：{exc}"
        ) from exc

    # 空页、空段落不写入向量库，否则会产生没有检索价值的空向量块。
    documents = [
        document
        for document in documents
        if document.page_content.strip()
    ]
    if not documents:
        raise DocumentParseError(
            f"文件 {original_name} 中没有提取到可用文本。"
        )

    return documents


def _load_pdf(
    file_path: Path,
    base_metadata: dict[str, str],
) -> list[Document]:
    """逐页解析 PDF，并标注页码和所属章节。"""

    with fitz.open(file_path) as pdf:
        if pdf.needs_pass:
            raise DocumentParseError("PDF 已加密，暂时无法解析。")

        # 先一次性读取各页文字，后续章节识别需要横向比较多个页面。
        page_texts = [page.get_text("text").strip() for page in pdf]

    chapter_ranges = _detect_pdf_chapter_ranges(page_texts)
    documents: list[Document] = []

    for page_index, text in enumerate(page_texts):
        if not text:
            continue

        page_number = page_index + 1
        metadata = dict(base_metadata)
        metadata["page_number"] = page_number

        # PDF 自带页码但通常没有结构化章节；这里按检测出的范围补充章节元数据。
        chapter = _find_page_chapter(page_number, chapter_ranges)
        if chapter is not None:
            chapter_number, chapter_title = chapter
            metadata["chapter_number"] = chapter_number
            metadata["chapter_title"] = chapter_title

        documents.append(
            Document(
                page_content=text,
                metadata=metadata,
            )
        )

    return documents


def _detect_pdf_chapter_ranges(
    page_texts: list[str],
) -> list[tuple[int, int, int, str]]:
    """根据重复页眉识别各章节在 PDF 中的实际页面范围。"""

    # 记录每章标题在哪些 PDF 页出现，后续用“重复且相邻”判断哪些是正文页眉。
    chapter_occurrences: dict[int, list[tuple[int, str]]] = {}

    for page_number, text in enumerate(page_texts, start=1):
        # set 去重：同一页的页眉和正文可能各出现一次相同章节标题。
        matches = {
            (
                _parse_chapter_number(raw_number),
                _clean_chapter_title(raw_title),
            )
            for raw_number, raw_title in CHAPTER_PATTERN.findall(text)
        }
        valid_matches = {
            (number, title)
            for number, title in matches
            if number is not None and title
        }

        # 目录页通常同时包含多个章节，不能作为章节正文起点。
        if len(valid_matches) != 1:
            continue

        chapter_number, chapter_title = next(iter(valid_matches))
        chapter_occurrences.setdefault(chapter_number, []).append(
            (page_number, chapter_title)
        )

    detected_chapters: list[tuple[int, int, int, str]] = []
    for chapter_number, occurrences in chapter_occurrences.items():
        # 连续出现次数最多的一组通常是正文；目录页只会产生孤立、短小的一组。
        clusters = _group_nearby_chapter_pages(occurrences)
        if not clusters:
            continue

        # 正文页眉会在多页重复，目录中的偶然单页会形成更短的分组。
        best_cluster = max(clusters, key=len)
        start_page = best_cluster[0][0]
        last_header_page = best_cluster[-1][0]
        title_counts = Counter(title for _, title in best_cluster)
        chapter_title = title_counts.most_common(1)[0][0]
        detected_chapters.append(
            (
                chapter_number,
                start_page,
                last_header_page,
                chapter_title,
            )
        )

    detected_chapters.sort(key=lambda item: item[1])
    ranges: list[tuple[int, int, int, str]] = []

    # 当前章节结束于下一章开始前一页；最后一章结束于最后一次检测到页眉的位置。
    for chapter_index, chapter in enumerate(detected_chapters):
        chapter_number, start_page, last_header_page, title = chapter
        if chapter_index + 1 < len(detected_chapters):
            end_page = detected_chapters[chapter_index + 1][1] - 1
        else:
            end_page = last_header_page

        ranges.append((chapter_number, start_page, end_page, title))

    return ranges


def _group_nearby_chapter_pages(
    occurrences: list[tuple[int, str]],
) -> list[list[tuple[int, str]]]:
    """把相邻页面中重复出现的同一章节页眉分成一组。"""

    if not occurrences:
        return []

    # 先按页码排序，再依据页码间隔切成多个连续簇。
    sorted_occurrences = sorted(occurrences)
    clusters = [[sorted_occurrences[0]]]

    for occurrence in sorted_occurrences[1:]:
        previous_page = clusters[-1][-1][0]
        if occurrence[0] - previous_page <= MAX_CHAPTER_PAGE_GAP:
            clusters[-1].append(occurrence)
        else:
            clusters.append([occurrence])

    return clusters


def _find_page_chapter(
    page_number: int,
    chapter_ranges: list[tuple[int, int, int, str]],
) -> tuple[int, str] | None:
    """返回指定 PDF 页所属的章节。"""

    for chapter_number, start_page, end_page, title in chapter_ranges:
        if start_page <= page_number <= end_page:
            return chapter_number, title

    return None


def _parse_chapter_number(value: str) -> int | None:
    """把阿拉伯数字或常见中文数字转换为整数。"""

    if value.isdigit():
        return int(value)

    digit_values = {
        "一": 1,
        "二": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
    }
    if value in digit_values:
        return digit_values[value]
    if value == "十":
        return 10
    if value.startswith("十") and value[1:] in digit_values:
        return 10 + digit_values[value[1:]]
    if value.endswith("十") and value[:-1] in digit_values:
        return digit_values[value[:-1]] * 10
    if "十" in value:
        tens, ones = value.split("十", maxsplit=1)
        if tens in digit_values and ones in digit_values:
            return digit_values[tens] * 10 + digit_values[ones]

    return None


def _clean_chapter_title(value: str) -> str:
    """清理章节标题后可能跟随的正文或目录点线。"""

    title = re.split(r"[…·]{2,}|\s{2,}", value, maxsplit=1)[0]
    return title.strip(" ：:。.;；、-")[:40]


def _load_docx(
    file_path: Path,
    base_metadata: dict[str, str],
) -> list[Document]:
    """按 Word 标题分组，保留章节名供引用展示。"""

    # python-docx 能保留段落样式，因此可利用 Heading 标题把 Word 内容分段。
    word_document = WordDocument(file_path)
    documents: list[Document] = []
    current_title = "正文"
    current_paragraphs: list[str] = []

    def flush_section() -> None:
        # 遇到下一个标题或文档结尾时，把累积的正文落成一个 Document。
        content = "\n".join(current_paragraphs).strip()
        if not content:
            return

        metadata = dict(base_metadata)
        metadata["section_title"] = current_title
        documents.append(Document(page_content=content, metadata=metadata))
        current_paragraphs.clear()

    for paragraph in word_document.paragraphs:
        paragraph_text = paragraph.text.strip()
        if not paragraph_text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            # 标题本身成为下一段的 section_title，不混入当前段的正文。
            flush_section()
            current_title = paragraph_text
            continue

        current_paragraphs.append(paragraph_text)

    # 表格内容经常承载重要信息，第一版将每行转换为制表符分隔文本。
    for table in word_document.tables:
        for row in table.rows:
            cell_values = [cell.text.strip() for cell in row.cells]
            row_text = "\t".join(value for value in cell_values if value)
            if row_text:
                current_paragraphs.append(row_text)

    flush_section()
    return documents


def _load_markdown(
    file_path: Path,
    base_metadata: dict[str, str],
) -> list[Document]:
    """按 Markdown 标题拆分成逻辑章节。"""

    # Markdown 不依赖额外解析库，按 # 标题切分即可保留最常用的结构信息。
    text = _read_text_with_fallback(file_path)
    documents: list[Document] = []
    current_title = "正文"
    current_lines: list[str] = []

    def flush_section() -> None:
        # 与 Word 解析相同：在标题边界将已收集的行输出为一个逻辑段落。
        content = "\n".join(current_lines).strip()
        if not content:
            return

        metadata = dict(base_metadata)
        metadata["section_title"] = current_title
        documents.append(Document(page_content=content, metadata=metadata))
        current_lines.clear()

    for line in text.splitlines():
        stripped_line = line.strip()
        if stripped_line.startswith("#"):
            heading_text = stripped_line.lstrip("#").strip()
            if heading_text:
                flush_section()
                current_title = heading_text
                continue

        current_lines.append(line)

    flush_section()
    return documents


def _load_text(
    file_path: Path,
    base_metadata: dict[str, str],
) -> list[Document]:
    """解析常见 UTF-8 或 GB18030 编码的文本。"""

    text = _read_text_with_fallback(file_path)
    return [Document(page_content=text, metadata=dict(base_metadata))]


def _read_text_with_fallback(file_path: Path) -> str:
    """依次尝试常见中文文本编码。"""

    # utf-8-sig 会自动去掉 Windows 常见的 BOM；GB18030 覆盖大多数中文旧编码文本。
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise DocumentParseError("无法识别文本文件编码。")

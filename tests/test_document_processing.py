"""文档解析和分块测试。"""

from pathlib import Path

import fitz
from docx import Document as WordDocument

from src.document_loaders import load_document, split_documents
from src.document_loaders.parser import _detect_pdf_chapter_ranges


def test_load_utf8_text_and_split(tmp_path: Path) -> None:
    """验证 UTF-8 文本会保留来源元数据，并被分成唯一 ID 的多个块。"""

    file_path = tmp_path / "knowledge.txt"
    file_path.write_text(
        "LangChain 可以连接文档检索与大语言模型。\n" * 20,
        encoding="utf-8",
    )

    documents = load_document(
        file_path,
        document_id="document-1",
        knowledge_base_id="kb-1",
        original_name="knowledge.txt",
    )
    chunks, chunk_ids = split_documents(
        documents,
        document_id="document-1",
        chunk_size=120,
        chunk_overlap=20,
    )

    assert documents[0].metadata["source_file"] == "knowledge.txt"
    assert len(chunks) > 1
    assert len(chunks) == len(chunk_ids)
    assert len(set(chunk_ids)) == len(chunk_ids)
    assert chunks[0].metadata["chunk_index"] == 0


def test_markdown_keeps_section_title(tmp_path: Path) -> None:
    """验证 Markdown 的一级标题会成为对应文本段的 section_title。"""

    file_path = tmp_path / "notes.md"
    file_path.write_text(
        "# 第一章\n这是第一章内容。\n\n# 第二章\n这是第二章内容。",
        encoding="utf-8",
    )

    documents = load_document(
        file_path,
        document_id="document-2",
        knowledge_base_id="kb-1",
        original_name="notes.md",
    )

    assert len(documents) == 2
    assert documents[0].metadata["section_title"] == "第一章"
    assert documents[1].metadata["section_title"] == "第二章"


def test_pdf_keeps_page_number(tmp_path: Path) -> None:
    """验证 PDF 逐页解析后，文本块仍记录原始页码。"""

    # 在临时目录现场生成最小 PDF，避免依赖仓库中的外部测试文件。
    file_path = tmp_path / "manual.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF page content")
    pdf.save(file_path)
    pdf.close()

    documents = load_document(
        file_path,
        document_id="document-pdf",
        knowledge_base_id="kb-1",
        original_name="manual.pdf",
    )

    assert len(documents) == 1
    assert documents[0].metadata["page_number"] == 1
    assert "PDF page content" in documents[0].page_content


def test_docx_keeps_heading(tmp_path: Path) -> None:
    """验证 Word Heading 样式会成为后续正文的章节标题。"""

    file_path = tmp_path / "manual.docx"
    word_document = WordDocument()
    word_document.add_heading("安装说明", level=1)
    word_document.add_paragraph("这是安装步骤。")
    word_document.save(file_path)

    documents = load_document(
        file_path,
        document_id="document-docx",
        knowledge_base_id="kb-1",
        original_name="manual.docx",
    )

    assert len(documents) == 1
    assert documents[0].metadata["section_title"] == "安装说明"
    assert "这是安装步骤" in documents[0].page_content


def test_pdf_chapter_ranges_ignore_toc_and_cover_chapter_pages() -> None:
    """验证目录页不被误判为正文，并正确计算每章实际页码范围。"""

    page_texts = [
        "目录\n第1章 有理数\n第2章 有理数的运算",
        "第1章 有理数\n章节首页",
        "第1章 有理数\n章节正文",
        "第2章 有理数的运算\n章节首页",
        "第2章 有理数的运算\n章节正文",
    ]

    ranges = _detect_pdf_chapter_ranges(page_texts)

    assert ranges == [
        (1, 2, 3, "有理数"),
        (2, 4, 5, "有理数的运算"),
    ]
